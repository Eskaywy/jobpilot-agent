"""Step 6 - Excel Tracking Ledger (pandas + openpyxl).

Maintains ``./tracker.xlsx``: creates the workbook with a styled header on
first use, appends one row per processed listing every cycle, and exposes the
known (company, role, recipient) triples so the pipeline never applies twice.

Recorded fields (spec): Timestamp, Company, Job Role, Recipient Email,
Keyword Match %, Resume Used, Cover Letter Generated, Application Status -
plus Cosine Similarity, Keyword Coverage, Source, Job URL and Notes for
full auditability of the matching decision.
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Set

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .textutils import dedup_key

log = logging.getLogger("jobpilot.tracker")

#: Canonical column order for tracker.xlsx.
COLUMNS: List[str] = [
    "Timestamp", "Company", "Job Role", "Recipient Email",
    "Keyword Match %", "Cosine Similarity", "Keyword Coverage",
    "Resume Used", "Cover Letter Generated", "Application Status",
    "Source", "Job URL", "Notes",
]

#: Column widths (chars) keyed by header name.
COLUMN_WIDTHS: Dict[str, int] = {
    "Timestamp": 20, "Company": 26, "Job Role": 30, "Recipient Email": 32,
    "Keyword Match %": 15, "Cosine Similarity": 16, "Keyword Coverage": 16,
    "Resume Used": 42, "Cover Letter Generated": 18,
    "Application Status": 18, "Source": 12, "Job URL": 34, "Notes": 34,
}


class Tracker:
    """Append-only Excel ledger with SHA-256 backed deduplication."""

    def __init__(self, path: Path):
        self.path = Path(path)

    # ---------------------------------------------------------------- setup
    def ensure(self) -> None:
        """Create the workbook with headers + styling if it doesn't exist."""
        if self.path.exists():
            return
        try:
            pd.DataFrame(columns=COLUMNS).to_excel(
                self.path, index=False, sheet_name="Applications")
            self._style()
            log.info("Created tracking ledger: %s", self.path)
        except OSError as exc:
            log.error("Cannot create tracker at %s: %s "
                      "(close it in Excel Viewer if open)", self.path, exc)

    def _style(self) -> None:
        """Bold white-on-navy header, frozen top row, sensible widths."""
        from openpyxl import load_workbook
        workbook = load_workbook(self.path)
        sheet = workbook["Applications"]
        fill = PatternFill("solid", fgColor="1F3864")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for index, column in enumerate(COLUMNS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                COLUMN_WIDTHS.get(column, 18))
        sheet.freeze_panes = "A2"
        workbook.save(self.path)

    # ----------------------------------------------------------------- read
    def known_keys(self) -> Set[str]:
        """All previously recorded (company, role, recipient) dedup keys."""
        if not self.path.exists():
            return set()
        try:
            frame = pd.read_excel(self.path, engine="openpyxl")
        except Exception as exc:
            log.warning("Could not read tracker (%s) - dedup disabled "
                        "for this cycle", exc)
            return set()
        frame = frame.fillna("")
        return {
            dedup_key(str(row.get("Company", "")),
                      str(row.get("Job Role", "")),
                      str(row.get("Recipient Email", "")))
            for _, row in frame.iterrows()
        }

    # ---------------------------------------------------------------- write
    def append(self, row: Dict[str, object]) -> None:
        """Append one application record and re-apply the styling."""
        try:
            frame = pd.read_excel(self.path, engine="openpyxl")
        except Exception:
            frame = pd.DataFrame(columns=COLUMNS)
        frame = pd.concat([frame, pd.DataFrame([row])], ignore_index=True)
        frame = frame.reindex(columns=COLUMNS)  # enforce canonical order
        try:
            frame.to_excel(self.path, index=False, sheet_name="Applications")
            self._style()
            log.info("Tracker updated: %s | %s | %s",
                     row.get("Company"), row.get("Job Role"),
                     row.get("Application Status"))
        except OSError as exc:
            # Most common cause: tracker.xlsx open in Excel Viewer.
            log.error("Failed to write tracker (%s). Row withheld: %s", exc, row)
